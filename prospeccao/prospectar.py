#!/usr/bin/env python3
"""Prospecção da Dri: rastreia clínicas, detecta reclamação de demora e envia email.

Uso:
  python3 prospectar.py rastrear          # busca clínicas no Google e preenche a planilha
  python3 prospectar.py enviar --teste    # mostra os emails que seriam enviados, sem enviar
  python3 prospectar.py enviar            # envia de verdade (limite por execução no config)
  python3 prospectar.py tudo              # rastrear + enviar

Variáveis de ambiente:
  GOOGLE_PLACES_API_KEY   chave da API Places (New) do Google Cloud
  GMAIL_APP_PASSWORD      senha de app do Gmail (myaccount.google.com/apppasswords)
"""
import argparse
import datetime as dt
import json
import os
import re
import smtplib
import sys
import time
import urllib.parse
import urllib.request
from email.message import EmailMessage
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

AQUI = Path(__file__).resolve().parent
ABA = "Prospecção"
LINHA_CABECALHO = 4
PRIMEIRA_LINHA = 5
LINHA_FUNIL_ORIGINAL = 205  # fórmulas do Funil cobrem até aqui; ampliamos se passar

# Colunas existentes da planilha (1 = A)
C_NUM, C_CLINICA, C_ESPEC, C_BAIRRO, C_WHATS, C_INSTA = 1, 2, 3, 4, 5, 6
C_ORIGEM, C_ABORDAGEM, C_STATUS, C_OBS = 10, 14, 18, 21
# Colunas novas, depois de "Observações"
NOVAS = ["Cidade", "Nota Google", "Nº de avaliações", "Reclamação de demora (trecho)",
         "Site", "Email", "Google Place ID", "Fonte do email"]
(C_CIDADE, C_NOTA, C_NAVAL, C_RECLAMA, C_SITE, C_EMAIL, C_PLACE,
 C_FONTE_EMAIL) = range(22, 22 + len(NOVAS))

PADRAO_DEMORA = re.compile(
    r"demor|esper(a|ei|ando|ar)\b|aguard|fila|atras|"
    r"n[ãa]o (me )?(atend|respond|retorn)|ningu[ée]m (atend|respond)|"
    r"sem (resposta|retorno)|dif[íi]cil (marcar|agendar|falar|contato)|"
    r"horas? (pra|para) ser atendid|nunca atendem|visualiza e n[ãa]o",
    re.IGNORECASE,
)
PADRAO_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
EMAIL_IGNORAR = re.compile(r"\.(png|jpe?g|gif|webp|svg)$|sentry|wixpress|example|exemplo|"
                           r"godaddy|seudominio|email@|nome@", re.IGNORECASE)


def carregar_config():
    caminho = AQUI / "config.json"
    if not caminho.exists():
        sys.exit("Crie prospeccao/config.json a partir de config.exemplo.json.")
    return json.loads(caminho.read_text(encoding="utf-8"))


# ---------------------------------------------------------------- planilha

def abrir_planilha(cfg):
    caminho = (AQUI / cfg["planilha"]).resolve()
    wb = openpyxl.load_workbook(caminho)
    ws = wb[ABA]
    negrito = ws.cell(LINHA_CABECALHO, C_OBS).font
    preenchimento = ws.cell(LINHA_CABECALHO, C_OBS).fill
    for i, nome in enumerate(NOVAS):
        cel = ws.cell(LINHA_CABECALHO, C_CIDADE + i)
        if cel.value != nome:
            cel.value = nome
            cel.font = Font(bold=negrito.bold, color=negrito.color)
            cel.fill = PatternFill(fill_type=preenchimento.fill_type,
                                   fgColor=preenchimento.fgColor)
    return wb, ws, caminho


def linhas_preenchidas(ws):
    for r in range(PRIMEIRA_LINHA, ws.max_row + 1):
        if ws.cell(r, C_CLINICA).value:
            yield r


def proxima_linha_vazia(ws):
    r = PRIMEIRA_LINHA
    while ws.cell(r, C_CLINICA).value:
        r += 1
    return r


def ampliar_funil(wb, ws):
    """Com a expansão para o Brasil a lista passa da linha 205: estende as fórmulas
    do Funil e as listas suspensas para cobrir todas as linhas."""
    ultima = max(LINHA_FUNIL_ORIGINAL, max(linhas_preenchidas(ws), default=0) + 200)
    padrao = re.compile(r"(Prospecção!\$?[A-Z]+\$?5:\$?[A-Z]+\$?)(\d+)")
    for linha in wb["Funil"].iter_rows():
        for cel in linha:
            if isinstance(cel.value, str) and cel.value.startswith("="):
                cel.value = padrao.sub(lambda m: m.group(1) + str(ultima), cel.value)
    for dv in ws.data_validations.dataValidation:
        col = str(dv.sqref).split(":")[0].rstrip("0123456789")
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"{col}5:{col}{ultima}")
    for r in range(PRIMEIRA_LINHA, ultima + 1):
        if ws.cell(r, C_NUM).value is None:
            ws.cell(r, C_NUM).value = r - PRIMEIRA_LINHA + 1


def normalizar(nome):
    return re.sub(r"\W+", "", (nome or "").lower())


# ---------------------------------------------------------------- rastrear

def buscar_places(chave, consulta, pagina=None):
    corpo = {"textQuery": consulta, "languageCode": "pt-BR", "regionCode": "BR"}
    if pagina:
        corpo["pageToken"] = pagina
    campos = ",".join([
        "places.id", "places.displayName", "places.formattedAddress",
        "places.addressComponents", "places.rating", "places.userRatingCount",
        "places.nationalPhoneNumber", "places.websiteUri", "places.reviews",
        "places.businessStatus", "nextPageToken",
    ])
    req = urllib.request.Request(
        "https://places.googleapis.com/v1/places:searchText",
        data=json.dumps(corpo).encode(),
        headers={"Content-Type": "application/json", "X-Goog-Api-Key": chave,
                 "X-Goog-FieldMask": campos},
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.load(resp)


def bairro_de(lugar):
    for comp in lugar.get("addressComponents", []):
        if "sublocality" in comp.get("types", []) or "sublocality_level_1" in comp.get("types", []):
            return comp.get("longText")
    return None


def trecho_demora(lugar):
    """Primeira avaliação de 1 a 3 estrelas que fala de demora, cortada em ~200 caracteres."""
    for av in lugar.get("reviews", []):
        texto = (av.get("originalText") or av.get("text") or {}).get("text", "")
        if av.get("rating", 5) <= 3 and PADRAO_DEMORA.search(texto):
            texto = " ".join(texto.split())
            return texto[:200] + ("…" if len(texto) > 200 else "")
    return None


PADRAO_CNPJ = re.compile(r"\b(\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2})\b")
PADRAO_INSTA = re.compile(r"instagram\.com/([A-Za-z0-9_.]{2,30})", re.IGNORECASE)
PADRAO_CFEMAIL = re.compile(r'data-cfemail="([0-9a-f]+)"')
PADRAO_LINK_CONTATO = re.compile(r'href="([^"#]*(?:contat|fale|atendimento|sobre|quem-somos)[^"#]*)"',
                                 re.IGNORECASE)
PAGINAS_CONTATO = ("", "contato", "fale-conosco", "contact", "sobre", "quem-somos")


def baixar(url, limite=800_000):
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            return resp.read(limite).decode("utf-8", "ignore")
    except Exception:
        return ""


def emails_no_html(html):
    # Cloudflare esconde o email em data-cfemail; o primeiro byte é a chave do XOR
    for cod in PADRAO_CFEMAIL.findall(html):
        chave = int(cod[:2], 16)
        html += " " + "".join(chr(int(cod[i:i + 2], 16) ^ chave) for i in range(2, len(cod), 2))
    html = re.sub(r"\s*(\[at\]|\(at\)|\[arroba\])\s*", "@", html, flags=re.IGNORECASE)
    return [e.lower() for e in PADRAO_EMAIL.findall(html) if not EMAIL_IGNORAR.search(e)]


def email_pelo_cnpj(cnpj):
    """Email cadastrado na Receita Federal (dado público). Às vezes é o do contador."""
    cnpj = re.sub(r"\D", "", cnpj)
    for url, caminho in ((f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}", ("email",)),
                         (f"https://publica.cnpj.ws/cnpj/{cnpj}", ("estabelecimento", "email"))):
        texto = baixar(url, 200_000)
        if not texto:
            continue
        try:
            dado = json.loads(texto)
            for chave in caminho:
                dado = dado.get(chave) if isinstance(dado, dict) else None
        except ValueError:
            continue
        if dado and PADRAO_EMAIL.fullmatch(dado.strip()):
            return dado.strip().lower()
    return None


def buscar_contato(site):
    """Procura email e Instagram no site da clínica. Ordem:
    1. páginas de contato do site (inclui email escondido pelo Cloudflare e "[arroba]");
    2. CNPJ do rodapé consultado na Receita (BrasilAPI / CNPJ.ws).
    Devolve (email, fonte do email, instagram)."""
    if not site:
        return None, None, None
    base = site if site.startswith("http") else "https://" + site
    base = base if base.endswith("/") else base + "/"
    dominio = urllib.parse.urlparse(base).netloc.replace("www.", "")
    paginas = [urllib.parse.urljoin(base, c) for c in PAGINAS_CONTATO]
    emails, cnpjs, insta, i = [], [], None, 0
    while i < len(paginas) and i < 10:
        html = baixar(paginas[i])
        if i == 0:  # links de contato que a home aponta
            paginas += [urllib.parse.urljoin(base, h) for h in PADRAO_LINK_CONTATO.findall(html)
                        if dominio in urllib.parse.urljoin(base, h)]
        emails += emails_no_html(html)
        cnpjs += PADRAO_CNPJ.findall(html)
        m = PADRAO_INSTA.search(html)
        if m and not insta and m.group(1).lower() not in ("p", "reel", "explore", "accounts"):
            insta = "@" + m.group(1).rstrip(".")
        i += 1
        if emails and insta:
            break
    if emails:
        emails.sort(key=lambda e: (dominio not in e, e))  # prefere o domínio da clínica
        return emails[0], "Site da clínica", insta
    for cnpj in dict.fromkeys(cnpjs):
        email = email_pelo_cnpj(cnpj)
        if email:
            return email, f"CNPJ {cnpj} (Receita)", insta
    return None, None, insta


def rastrear(cfg):
    chave = os.environ.get("GOOGLE_PLACES_API_KEY")
    if not chave:
        sys.exit("Defina GOOGLE_PLACES_API_KEY.")
    wb, ws, caminho = abrir_planilha(cfg)
    vistos = {normalizar(ws.cell(r, C_CLINICA).value) for r in linhas_preenchidas(ws)}
    vistos |= {ws.cell(r, C_PLACE).value for r in linhas_preenchidas(ws)}

    # Cada cidade pode ter bairros (busca mais fina) ou não (busca na cidade toda)
    alvos = [(c["cidade"], b) for c in cfg["cidades"] for b in (c.get("bairros") or [None])]
    novos = []
    for termo in cfg["buscas"]:
        for cidade, bairro in alvos:
            consulta = " ".join(x for x in (termo, bairro, cidade) if x)
            pagina = None
            for _ in range(3):  # até 60 resultados por consulta
                try:
                    dados = buscar_places(chave, consulta, pagina)
                except Exception as e:
                    print(f"  erro em '{consulta}': {e}")
                    break
                for lugar in dados.get("places", []):
                    nome = lugar.get("displayName", {}).get("text")
                    if (not nome or lugar["id"] in vistos or normalizar(nome) in vistos
                            or lugar.get("businessStatus") not in (None, "OPERATIONAL")):
                        continue
                    vistos |= {lugar["id"], normalizar(nome)}
                    trecho = trecho_demora(lugar)
                    if not trecho:
                        continue  # só interessam as que têm reclamação de demora
                    novos.append({
                        "nome": nome, "especialidade": termo.replace("clínica ", "").capitalize(),
                        "bairro": bairro_de(lugar) or bairro, "cidade": cidade,
                        "telefone": lugar.get("nationalPhoneNumber"),
                        "nota": lugar.get("rating"), "n": lugar.get("userRatingCount"),
                        "trecho": trecho, "site": lugar.get("websiteUri"), "place": lugar["id"],
                    })
                    print(f"  + {nome} ({lugar.get('rating')}★)")
                pagina = dados.get("nextPageToken")
                if not pagina:
                    break
                time.sleep(2)

    print(f"{len(novos)} clínicas novas com reclamação de demora. Procurando emails…")
    for c in novos:
        c["email"], c["fonte_email"], c["insta"] = buscar_contato(c["site"])

    # Quem tem email e mais avaliações primeiro
    novos.sort(key=lambda c: (c["email"] is None, -(c["n"] or 0)))
    gravados = 0
    for c in novos:
        r = proxima_linha_vazia(ws)
        valores = {C_CLINICA: c["nome"], C_ESPEC: c["especialidade"], C_BAIRRO: c["bairro"],
                   C_WHATS: c["telefone"], C_ORIGEM: "Outro", C_STATUS: "Na lista",
                   C_OBS: "Rastreio automático (Google)", C_NOTA: c["nota"], C_NAVAL: c["n"],
                   C_RECLAMA: c["trecho"], C_SITE: c["site"], C_EMAIL: c["email"],
                   C_PLACE: c["place"], C_CIDADE: c["cidade"],
                   C_FONTE_EMAIL: c["fonte_email"], C_INSTA: c["insta"]}
        for col, v in valores.items():
            ws.cell(r, col).value = v
        gravados += 1
    ampliar_funil(wb, ws)
    wb.save(caminho)
    print(f"{gravados} clínicas gravadas em {caminho.name}.")


# ---------------------------------------------------------------- enviar

def montar_email(cfg, clinica):
    assunto = f"Pacientes da {clinica} esperando resposta no WhatsApp"
    corpo = f"""Olá, equipe da {clinica},

Sou o {cfg['remetente_nome']}, enfermeiro hospitalar há 15 anos e criador da Dri. Olhando as avaliações públicas da {clinica}, vi que alguns pacientes comentam sobre demora para conseguir resposta ou atendimento. Não é crítica: acontece em quase toda clínica com muita procura, e o paciente que não recebe retorno rápido costuma marcar com quem responde primeiro.

A Dri é uma atendente de IA no WhatsApp que responde em segundos, 24 horas por dia, e marca a consulta direto no horário real de cada médico.

Para ver como funciona, é só clicar no link e conversar com ela como se fosse paciente: {cfg['link_demonstracao']}

Se fizer sentido para a {clinica}, me responda aqui ou no WhatsApp que eu já envio o contrato e deixo a Dri funcionando com a agenda de vocês.

Abraço,
{cfg['remetente_nome']}
WhatsApp: {cfg['assinatura_whatsapp']}

Se não quiser receber outras mensagens minhas, é só responder "remover".
"""
    return assunto, corpo


def enviar(cfg, teste):
    wb, ws, caminho = abrir_planilha(cfg)
    fila = [r for r in linhas_preenchidas(ws)
            if ws.cell(r, C_EMAIL).value and not ws.cell(r, C_ABORDAGEM).value
            and ws.cell(r, C_STATUS).value in (None, "Na lista", "Teste oculto feito")]
    fila = fila[: cfg["limite_envios_por_execucao"]]
    if not fila:
        print("Nenhuma clínica com email aguardando abordagem.")
        return

    smtp = None
    if not teste:
        senha = os.environ.get("GMAIL_APP_PASSWORD")
        if not senha:
            sys.exit("Defina GMAIL_APP_PASSWORD.")
        smtp = smtplib.SMTP_SSL("smtp.gmail.com", 465)
        smtp.login(cfg["remetente_email"], senha)

    hoje = dt.datetime.combine(dt.date.today(), dt.time())
    for i, r in enumerate(fila):
        clinica = ws.cell(r, C_CLINICA).value
        destino = ws.cell(r, C_EMAIL).value
        assunto, corpo = montar_email(cfg, clinica)
        if teste:
            print(f"--- PARA: {destino}\nASSUNTO: {assunto}\n\n{corpo}")
            continue
        msg = EmailMessage()
        msg["From"] = f"{cfg['remetente_nome']} <{cfg['remetente_email']}>"
        msg["To"] = destino
        msg["Subject"] = assunto
        msg.set_content(corpo)
        try:
            smtp.send_message(msg)
        except Exception as e:
            print(f"  falhou {clinica} <{destino}>: {e}")
            continue
        ws.cell(r, C_ABORDAGEM).value = hoje
        ws.cell(r, C_ABORDAGEM).number_format = "DD/MM/YYYY"
        ws.cell(r, C_STATUS).value = "Abordado"
        ws.cell(r, 19).value = "Follow-up se não responder"  # Próxima ação
        ws.cell(r, 20).value = hoje + dt.timedelta(days=3)  # Data da próxima ação
        ws.cell(r, 20).number_format = "DD/MM/YYYY"
        wb.save(caminho)  # salva a cada envio para não reenviar se cair no meio
        print(f"  enviado: {clinica} <{destino}>")
        if i < len(fila) - 1:
            time.sleep(cfg["intervalo_entre_envios_segundos"])
    if smtp:
        smtp.quit()


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("acao", choices=["rastrear", "enviar", "tudo"])
    p.add_argument("--teste", action="store_true", help="não envia, só mostra os emails")
    a = p.parse_args()
    cfg = carregar_config()
    if a.acao in ("rastrear", "tudo"):
        rastrear(cfg)
    if a.acao in ("enviar", "tudo"):
        enviar(cfg, a.teste)


if __name__ == "__main__":
    main()
